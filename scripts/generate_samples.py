from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "samples"


def format_header(sheet, row: int) -> None:
    for cell in sheet[row]:
        if cell.value is not None:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="26352B")


def write_shifted_headers() -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rates"
    sheet.append(["Northstar Capital | indicative deposit rates"])
    sheet.append(["Effective date: 2025-02-01", None, None, None, None])
    sheet.append(["Offering", "Term", "CCY", "Annual yield (%)", "Pricing date", "Notes"])
    sheet.append(["Flexible Reserve", "30 days", "USD", 4.10, "2025-02-01", "Online only"])
    sheet.append(["Flexible Reserve", "90 days", "USD", 4.40, "2025-02-01", ""])
    sheet.append(["Community Bond", "6 months", "GBP", 3.65, "2025-02-01", "New money"])
    format_header(sheet, 3)
    sheet.freeze_panes = "A4"
    workbook.save(SAMPLES / "shifted-headers.xlsx")


def write_mixed_sheets() -> None:
    workbook = Workbook()
    deposits = workbook.active
    deposits.title = "Term deposits"
    deposits.append(["Product name", "Maturity", "Currency", "Interest rate", "As of"])
    deposits.append(["Steady Term", "90 days", "CAD", 3.95, "2025-02-10"])
    deposits.append(["Steady Term", "1 year", "CAD", 4.20, "2025-02-10"])
    deposits.append(["Reserve Plus", "2 years", "USD", 4.35, "2025-02-10"])
    format_header(deposits, 1)
    deposits.freeze_panes = "A2"

    lending = workbook.create_sheet("Lending")
    lending.append(["Loan product", "Fixed period", "Rate (%)", "Denomination", "Pricing date"])
    lending.append(["Home Flex", "2 years", 5.15, "USD", "2025-02-10"])
    lending.append(["Home Flex", "5 years", 5.45, "USD", "2025-02-10"])
    lending.append(["Small Business", "3 years", 6.10, "EUR", "2025-02-10"])
    format_header(lending, 1)
    lending.freeze_panes = "A2"

    notes = workbook.create_sheet("Read me")
    notes.append(["Illustrative demo workbook"])
    notes.append(["Contains separate deposit and lending tables."])

    workbook.save(SAMPLES / "mixed-sheets.xlsx")


if __name__ == "__main__":
    SAMPLES.mkdir(parents=True, exist_ok=True)
    write_shifted_headers()
    write_mixed_sheets()
